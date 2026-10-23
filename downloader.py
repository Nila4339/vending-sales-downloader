"""JUSTECH report downloader prototype; site-specific selectors require live verification."""
import argparse
import csv
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

BASE = Path(__file__).resolve().parent

def safe_name(value):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip(" .") or "未命名"

def machines_from_csv():
    with (BASE / "machines.csv").open(encoding="utf-8-sig", newline="") as stream:
        items = list(csv.DictReader(stream))
    return [m for m in items if m.get("enabled", "").strip().lower() in ("1", "true", "yes", "y")]

def click(page, selector):
    page.locator(selector).first.click(timeout=20000)

def fill(page, selector, value):
    page.locator(selector).first.fill(value)

def validate_file(path, original_extension):
    if original_extension == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if not wb.sheetnames:
                raise ValueError("Excel 沒有工作表")
        finally:
            wb.close()
    elif original_extension == ".csv":
        if path.stat().st_size == 0:
            raise ValueError("CSV 是空檔案")
    elif original_extension == ".xls":
        if path.stat().st_size < 100:
            raise ValueError("XLS 檔案太小，需人工確認")
    else:
        raise ValueError("非預期匯出格式")

def save_log(path, row):
    existed = path.exists()
    with path.open("a", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=[
            "time", "machine_id", "machine_name", "billing_month", "status", "file", "error"
        ])
        if not existed:
            writer.writeheader()
        writer.writerow(row)

def export_one(page, config, machine, destination):
    """Site adapter stub. Requires live DOM selectors and fresh-export correlation.

    Do not treat the first notification as the current report without validation.
    """
    selectors = config["selectors"]
    required = (
        "machine_dropdown", "machine_option_template", "start_input", "end_input",
        "export_button", "report_dropdown", "report_option", "export_confirm",
        "notification_button", "notification_download_link",
    )
    missing = [key for key in required if not selectors.get(key)]
    if missing:
        raise RuntimeError("網站定位尚未完成：" + ", ".join(missing))

    click(page, selectors["machine_dropdown"])
    if selectors.get("machine_search"):
        fill(page, selectors["machine_search"], machine["machine_id"])
    machine_selector = selectors["machine_option_template"].replace(
        "{machine_id}", machine["machine_id"]
    ).replace("{machine_name}", machine["machine_name"])
    click(page, machine_selector)
    if selectors.get("machine_confirm"):
        click(page, selectors["machine_confirm"])

    # This direct-input method is only valid if the live site accepts keyboard-filled dates.
    fill(page, selectors["start_input"], config["start_datetime"])
    fill(page, selectors["end_input"], config["end_datetime"])
    if selectors.get("date_confirm"):
        click(page, selectors["date_confirm"])
    if selectors.get("search_button"):
        click(page, selectors["search_button"])

    # Safety gate: not yet possible to correlate a generated notification with
    # the current machine/date/job. Do not initiate an ambiguous real export.
    raise RuntimeError(
        "停止於安全檢查：通知中心的本次匯出任務辨識尚未驗證，"
        "需完成網站專屬 adapter 後才能正式匯出下載。"
    )

def main():
    parser = argparse.ArgumentParser(description="銷售報表逐台下載器 v0.1")
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    cfg_file = BASE / args.config
    if not cfg_file.exists():
        raise SystemExit("請先將 config.example.json 複製為 config.json")
    cfg = json.loads(cfg_file.read_text(encoding="utf-8"))
    machines = machines_from_csv()
    if not machines:
        raise SystemExit("machines.csv 尚無 enabled=1 機台")
    output = BASE / cfg["download_dir"] / cfg["billing_month"]
    for machine in machines:
        filename = f"{safe_name(cfg['billing_month'])}_{safe_name(machine['machine_name'])}_{safe_name(machine['machine_id'])}.xlsx"
        print(machine["machine_id"], "→", filename)
    if args.dry_run:
        return
    output.mkdir(parents=True, exist_ok=True)
    if not cfg["site_url"].startswith("https://") or "請填" in cfg["site_url"]:
        raise SystemExit("請在 config.json 填入有效 HTTPS 訂單頁網址")
    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(BASE / cfg["browser_profile_dir"]),
            headless=cfg.get("headless", False),
            accept_downloads=True,
        )
        try:
            page = context.pages[0] if context.pages else context.new_page()
            page.goto(cfg["site_url"])
            if cfg.get("manual_login", True):
                input("請手動登入並開啟訂單頁，再按 Enter：")
            for machine in machines:
                log = dict(
                    time=datetime.now().isoformat(timespec="seconds"),
                    machine_id=machine["machine_id"],
                    machine_name=machine["machine_name"],
                    billing_month=cfg["billing_month"],
                    status="", file="", error="",
                )
                try:
                    page.goto(cfg["site_url"])
                    path = export_one(page, cfg, machine, output)
                    log.update(status="success", file=str(path))
                except Exception as exc:
                    log.update(status="failed", error=str(exc))
                    print("失敗：", machine["machine_id"], exc)
                save_log(output / "download_log.csv", log)
        finally:
            context.close()

if __name__ == "__main__":
    main()
