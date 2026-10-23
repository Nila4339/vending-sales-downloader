"""v0.5.1 single-machine unattended-after-login validation.

Requires an initial Enter after login, but no per-export YES prompts.
Notification timestamp is NOT proof of device ownership; do not use while
other users are exporting the same report type concurrently.
"""
import argparse
import json
import re
from pathlib import Path
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeout
from test_date_range import desired_range, choose_day, set_clock

ROOT = Path(__file__).resolve().parent
URL = "https://lvmm.justech.com.tw/admin/19/orders"

def filename_safe(s):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", s).strip(" .") or "未命名"

def scan_notification_cards(page, report_pattern, submitted_at):
    """One browser roundtrip for all notification cards; no download URLs returned."""
    records = page.get_by_role("link", name="下載報表").evaluate_all("""links => links.slice(0,30).map((a, index) => {
        const text=(a.parentElement?.parentElement?.innerText||'').replace(/\\s+/g,' ').trim();
        return {index,text};
    })""")
    matching = []
    newest = None
    for record in records:
        match = report_pattern.search(record["text"])
        if not match:
            continue
        stamp = datetime.strptime(match.group(1)+match.group(2), "%Y%m%d%H%M%S").replace(tzinfo=ZoneInfo("Asia/Taipei"))
        newest = max(newest, stamp) if newest else stamp
        if (submitted_at - timedelta(seconds=2) <= stamp <=
                datetime.now(ZoneInfo("Asia/Taipei")) + timedelta(seconds=3)
                and "匯出成功" in record["text"]):
            matching.append((record["index"], match.group(0), stamp))
    return len(records), matching, newest


def diagnose_notification_cards(page, report_pattern, submitted_at):
    """Explain the first available card; do not expose download URLs."""
    rows = page.get_by_role("link", name="下載報表").evaluate_all("""links => links.slice(0, 3).map(a =>
      (a.parentElement?.parentElement?.innerText||'').replace(/\\s+/g, ' ').trim()
    )""")
    lower = submitted_at - timedelta(seconds=2)
    upper = datetime.now(ZoneInfo("Asia/Taipei")) + timedelta(seconds=3)
    print(f"  時間範圍：{lower:%H:%M:%S} 至 {upper:%H:%M:%S}")
    if not rows:
        print("  通知面板尚無可讀取的下載連結")
        return
    for i, row in enumerate(rows, 1):
        match = report_pattern.search(row)
        if not match:
            print(f"  通知 {i}：報表名稱格式未符合預期；內容摘要：{row[:120]}")
            continue
        stamp = datetime.strptime(match.group(1)+match.group(2), "%Y%m%d%H%M%S").replace(tzinfo=ZoneInfo("Asia/Taipei"))
        problems = []
        if not lower <= stamp <= upper:
            problems.append("時間不在範圍內")
        if "匯出成功" not in row:
            problems.append("沒有匯出成功標記")
        print(f"  通知 {i}：{match.group(0)}；原因：{', '.join(problems) if problems else '符合條件'}")


def reopen_notifications(page):
    """Reload orders page and reopen notifications, avoiding the covered toggle.

    The Filament notification overlay can intercept all clicks, even after Escape.
    Navigation clears it; the next notification open is a fresh UI request.
    """
    page.goto(URL, wait_until="domcontentloaded", timeout=30000)
    if "/login" in page.url:
        raise RuntimeError("登入狀態已失效，停止通知檢查")
    toggle = page.get_by_role("button", name="Open notifications")
    toggle.click(timeout=12000)
    page.wait_for_timeout(1800)



def clear_downloaded_notification(page, report_filename):
    """After verified download, dismiss only the matching card via its X.

    Filament cards use an icon-only close action. Match the exact report filename
    within the same card containing the download link, and never click global
    '清除' or other cards' controls.
    """
    links = page.get_by_role("link", name="下載報表")
    print(f"通知清理：重新開啟後可見 {links.count()} 筆下載通知，尋找 {report_filename}")
    matched = []
    for i in range(links.count()):
        card = links.nth(i).locator("xpath=../..")
        text = " ".join(card.inner_text().split())
        found = re.search(r"商品月銷售數量統計表_\d{8}_\d{6}\.xlsx", text)
        if found and found.group(0) == report_filename:
            matched.append(card)
    if len(matched) != 1:
        print(f"通知清除略過：對應卡片數 {len(matched)}，無法唯一定位。")
        return False
    card = matched[0]
    # Confirmed by read-only DOM inspection (2026-10-23):
    # download link -> parent^3: a single notification row with one icon-only
    # BUTTON containing an SVG. parent^6 contains multiple notifications.
    # Never use the global list "清除" control.
    single_row = card.locator("xpath=..")
    # 'card' is parent^2, so this row is parent^3.
    if single_row.get_by_role("link", name="下載報表").count() != 1:
        print("通知清除略過：單筆卡片含有非唯一下載連結。")
        return False
    row_text = " ".join(single_row.inner_text().split())
    match = re.search(r"商品月銷售數量統計表_\d{8}_\d{6}\.xlsx", row_text)
    if not match or match.group(0) != report_filename:
        print("通知清除略過：單筆卡片的報表名稱不符。")
        return False
    buttons = single_row.locator("button")
    if buttons.count() != 1:
        print(f"通知清除略過：單筆卡片有 {buttons.count()} 個按鈕，不符合實測結構。")
        return False
    close_button = buttons.first
    if close_button.locator("svg").count() != 1:
        print("通知清除略過：找不到預期的單個 SVG 關閉圖示。")
        return False
    print("通知清理：已找到本次報表卡片中的唯一 X，準備清除。")
    close_button.click(timeout=8000)
    page.wait_for_timeout(600)
    remaining = 0
    for i in range(page.get_by_role("link", name="下載報表").count()):
        text = " ".join(page.get_by_role("link", name="下載報表").nth(i).locator("xpath=../..").inner_text().split())
        found = re.search(r"商品月銷售數量統計表_\d{8}_\d{6}\.xlsx", text)
        if found and found.group(0) == report_filename:
            remaining += 1
    if remaining:
        print("已點擊單筆 X，但通知仍存在，請人工確認；不重複點擊。")
        return False
    print("已清除本次下載通知：", report_filename)
    return True


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--month", default="2026-10", help="對帳年月 YYYY-MM")
    parser.add_argument("--device-index", type=int, help="1-based index in device_names.json; skips interactive device selection")
    parser.add_argument("--skip-login-prompt", action="store_true", help="Batch mode: require pre-existing logged-in session")
    parser.add_argument("--result-file", help="Write exact downloaded filename for batch verification")
    args = parser.parse_args()
    names_path = ROOT / "device_names.json"
    if not names_path.exists():
        raise SystemExit("缺少 device_names.json，請先執行 discover_devices.py")
    names = json.loads(names_path.read_text(encoding="utf-8"))
    for idx, name in enumerate(names, 1):
        print(f"{idx:>3}. {name}")
    choice = str(args.device_index) if args.device_index is not None else input("請輸入要測試的設備序號：").strip()
    if not choice.isdigit() or not (1 <= int(choice) <= len(names)):
        raise SystemExit("設備序號不正確")
    name = names[int(choice)-1]
    start, end = desired_range(args.month)
    expected = start.strftime("%d/%m/%Y 12:00:00 AM") + " - " + end.strftime("%d/%m/%Y 12:00:00 AM")

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(ROOT / "browser_profile"), headless=False,
            viewport={"width":1440,"height":900}, accept_downloads=True)
        try:
            page = context.pages[0] if context.pages else context.new_page()
            page.goto(URL, wait_until="domcontentloaded")
            if not args.skip_login_prompt:
                input("請登入 JUSTECH、確認訂單頁已開啟，再按 Enter：")
                if "/login" in page.url:
                    page.goto(URL, wait_until="domcontentloaded")
            else:
                if "/login" in page.url:
                    raise RuntimeError("登入已失效；停止批次，請先執行單台模式登入")
                page.get_by_role("textbox", name="建立時間").wait_for(state="visible", timeout=15000)
            # Notification baseline was unreliable (zero even with historical messages).
            # Do not inspect old notifications here; keep the orders page untouched.
            print("通知中心將在匯出後檢查最新一筆；不再以舊連結數量為依據。")
            page.locator(".choices__inner").first.click()
            dropdown = page.locator(".choices__list--dropdown").first
            dropdown.wait_for(state="visible")
            search = dropdown.locator("input.choices__input")
            if search.count():
                search.first.fill(name)
            options = dropdown.locator(".choices__item--choice")
            matches = [i for i in range(options.count()) if " ".join(options.nth(i).inner_text().split()) == name]
            if len(matches) != 1:
                raise RuntimeError(f"設備選項無法唯一匹配 {name}：{len(matches)} 筆")
            options.nth(matches[0]).click()
            field = page.get_by_role("textbox",name="建立時間")
            field.click()
            picker = page.locator(".daterangepicker:visible").first
            choose_day(picker,start)
            choose_day(picker,end)
            set_clock(picker,"left","12","0","0","AM")
            set_clock(picker,"right","12","0","0","AM")
            picker.get_by_role("button",name="套用").click()
            actual = field.input_value().strip()
            if actual != expected:
                raise RuntimeError(f"日期不符：預期 {expected}，實際 {actual}")
            print("已選取設備：", name)
            print("已套用日期：", actual)
            print("自動檢查通過：機台選項唯一、日期區間正確，準備匯出。")

            page.get_by_role("button", name="匯出報表").click()
            # The recorded steps did not include choosing the export report type.
            # Do not silently assume a default selection.
            # Wait for the export dialog to render. Previous '.fi-modal-window:visible'
            # matched zero when the modal's actual visible container differed.
            page.get_by_text("確定要進行嗎？").wait_for(state="visible", timeout=15000)
            modal = page.locator(".fi-modal-window").filter(
                has=page.get_by_text("確定要進行嗎？")
            ).first
            if modal.count() != 1:
                raise RuntimeError("無法唯一辨識匯出確認視窗，停止")
            selectors = modal.locator(".fi-modal-content .choices__inner")
            try:
                selectors.first.wait_for(state="visible", timeout=10000)
            except Exception:
                print("診斷：匯出視窗數：", page.locator(".fi-modal-window").count())
                print("診斷：整頁 Choices 數：", page.locator(".choices__inner").count())
                print("診斷：此視窗 Choices 數：", selectors.count())
                raise RuntimeError("匯出視窗找不到可見的報表選單；請提供診斷數字與視窗截圖")
            if selectors.count() != 1:
                raise RuntimeError(f"匯出視窗有 {selectors.count()} 個報表選單，無法唯一定位")
            selectors.click()
            menu = modal.locator(".choices__list--dropdown").filter(
                has=page.locator(".choices__item--choice")
            ).first
            menu.wait_for(state="visible", timeout=10000)
            entries = menu.locator(".choices__item--choice")
            options = [" ".join(item.inner_text().split()) for item in entries.all()]
            print("匯出視窗的報表類型選項：", options)
            desired = "商品月銷售數量統計表"
            matches = [i for i, label in enumerate(options) if label == desired]
            if len(matches) != 1:
                raise RuntimeError("無法唯一找到指定報表，已停止。請提供上方選項。")
            entries.nth(matches[0]).click()
            print("已自動選取報表類型：", desired)
            # Choices.js sometimes leaves the dropdown overlay visible. Close it
            # by clicking the modal heading rather than forcing the submit button.
            if menu.is_visible():
                modal.get_by_text("確定要進行嗎？").click()
            try:
                menu.wait_for(state="hidden", timeout=5000)
            except Exception as exc:
                raise RuntimeError("報表選單未收合，為避免誤操作，停止送出") from exc
            chosen = " ".join(selectors.first.inner_text().split())
            if desired not in chosen:
                raise RuntimeError(f"報表選擇結果不符：{chosen!r}，停止送出")
            print("已確認選單收合，且選定報表正確。")
            print("自動檢查通過：報表類型正確，準備提交。")
            submitted_at = datetime.now(ZoneInfo("Asia/Taipei"))
            modal.get_by_role("button", name="確定").click(timeout=10000)
            print("本次提交起始時間（台灣）：", submitted_at.strftime("%Y-%m-%d %H:%M:%S"))
            print("已提交單台匯出，等候通知中心的新下載連結（最長 3 分鐘）...")
            page.wait_for_timeout(5000)
            print("匯出後已等待 5 秒，開始第一次通知檢查。")
            page.get_by_role("button", name="Open notifications").click()
            page.wait_for_timeout(1800)
            # v0.5: the tested notification card is two ancestors above its
            # download link; this card contains the timestamped report name.
            import time
            report_pattern = re.compile(r"商品月銷售數量統計表_(\d{8})_(\d{6})\.xlsx")
            deadline = time.monotonic() + 180
            selected_link = None
            selected_filename = None
            last_report = None
            attempts = 0
            while time.monotonic() < deadline:
                attempts += 1
                if attempts > 1:
                    reopen_notifications(page)
                count, candidates, newest = scan_notification_cards(page, report_pattern, submitted_at)
                newest_text = newest.strftime("%H:%M:%S") if newest else "無"
                print(f"通知檢查第 {attempts} 次：共 {count} 筆，最新報表時間 {newest_text}，符合 {len(candidates)} 筆", flush=True)
                if attempts == 1 or (attempts % 3 == 0 and len(candidates) == 0):
                    diagnose_notification_cards(page, report_pattern, submitted_at)
                if len(candidates) > 1:
                    print("時間範圍內的候選通知：",
                          [(filename, dt.strftime("%H:%M:%S")) for _,filename,dt in candidates])
                    raise RuntimeError("多筆通知時間吻合，無法辨識本次任務，停止下載")
                if len(candidates) == 1:
                    selected_index, selected_filename, stamp = candidates[0]
                    print(f"找到候選報表：{selected_filename}（匯出時間 {stamp:%H:%M:%S}）")
                    break
                page.wait_for_timeout(6000)
            if selected_filename is None:
                raise RuntimeError("3 分鐘內沒有找到本次提交時間區間內的唯一報表通知")
            page.wait_for_timeout(3000)
            reopen_notifications(page)
            count, settling_matches, newest = scan_notification_cards(page, report_pattern, submitted_at)
            if len(settling_matches) != 1 or settling_matches[0][1] != selected_filename:
                raise RuntimeError("通知再次核對後無法唯一匹配，停止自動下載")
            selected_index = settling_matches[0][0]
            selected_link = page.get_by_role("link", name="下載報表").nth(selected_index)
            print("通知類型、檔名及時間皆符合，準備自動下載；請勿在其他人同時匯出時使用。")
            # Recorder shows a popup and a download; popup is optional in runtime.
            pages_before_download = set(context.pages)
            with page.expect_download(timeout=60000) as download_info:
                selected_link.click()
            download = download_info.value
            # JUSTECH sometimes returns a UUID filename without any extension.
            # Inspect actual bytes before deciding the extension.
            import tempfile
            import zipfile
            import csv
            with tempfile.TemporaryDirectory() as tmpdir:
                temp = Path(tmpdir) / "downloaded_report"
                download.save_as(temp)
                if temp.stat().st_size == 0:
                    raise RuntimeError("下載內容為空")
                with temp.open("rb") as file_handle:
                    header = file_handle.read(8)
                if zipfile.is_zipfile(temp):
                    from openpyxl import load_workbook
                    try:
                        with temp.open("rb") as file_handle:
                            book = load_workbook(file_handle, read_only=True)
                            try:
                                if not book.sheetnames:
                                    raise ValueError("Excel 沒有工作表")
                            finally:
                                book.close()
                    except Exception as exc:
                        raise RuntimeError("下載檔是 ZIP，但無法辨識成 Excel xlsx") from exc
                    extension = ".xlsx"
                elif header.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                    extension = ".xls"
                else:
                    try:
                        sample = temp.read_bytes()[:4096].decode("utf-8-sig")
                        if not sample.strip() or sample.lstrip().lower().startswith(("<html", "<!doctype")):
                            raise ValueError("不是 CSV 內容")
                        rows = list(csv.reader(sample.splitlines()))
                        if len(rows) == 0 or max(map(len, rows)) < 2:
                            raise ValueError("未辨識出 CSV 欄位")
                        extension = ".csv"
                    except Exception as exc:
                        raise RuntimeError(
                            f"檔案格式無法辨識（原始檔名 {download.suggested_filename}），"
                            "先保留人工檢查，不套用 Excel 副檔名"
                        ) from exc
                target_dir = ROOT / "downloads" / args.month
                target_dir.mkdir(parents=True, exist_ok=True)
                target = target_dir / f"{filename_safe(args.month)}_{filename_safe(name)}{extension}"
                if target.exists():
                    stem = target.stem
                    suffix = target.suffix
                    counter = 2
                    while (target_dir / f"{stem}_第{counter}次{suffix}").exists():
                        counter += 1
                    target = target_dir / f"{stem}_第{counter}次{suffix}"
                    print("原檔案已存在，為避免覆蓋，改存為：", target.name)
                target.write_bytes(temp.read_bytes())
            print("下載完成：",target)
            if args.result_file:
                result_path = Path(args.result_file)
                result_path.write_text(json.dumps({"month":args.month,"device_index":int(choice),"device_name":name,"file":str(target.resolve()),"size":target.stat().st_size},ensure_ascii=False),encoding="utf-8")
            # The file is now validated and stored. Never clear other reports.
            try:
                reopen_notifications(page)
                clear_downloaded_notification(page, selected_filename)
            except Exception as exc:
                print("通知清除失敗，不影響下載結果：", exc)
            # JUSTECH may open a new browser tab when the download link is clicked.
            # Close only tabs created during this download, preserving the original orders tab.
            for extra_page in list(context.pages):
                if extra_page is not page and extra_page not in pages_before_download:
                    try:
                        extra_page.close()
                    except Exception as exc:
                        print("關閉下載分頁時出現提醒：", exc)
            if page.is_closed():
                raise RuntimeError("原訂單分頁已關閉，無法安全繼續")
            # Return to the known clean orders route. This avoids leftover notification
            # panels and stale machine/date filters when transitioning between machines.
            page.goto(URL, wait_until="domcontentloaded", timeout=30000)
            if "/login" in page.url:
                raise RuntimeError("登入狀態已失效；請重新登入後再執行")
            page.get_by_role("textbox", name="建立時間").wait_for(timeout=15000)
            page.bring_to_front()
            print("已關閉額外下載分頁，返回原訂單頁。")
            print("v0.5.1 單台無 YES 驗證完成；尚未驗證多台無人值守，勿直接大量執行。")
        finally:
            context.close()

if __name__ == "__main__":
    main()
