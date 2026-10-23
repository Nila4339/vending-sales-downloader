"""Single-device assisted download test. Selects device and date automatically.

You manually click export + notification download once. The script captures the
download event and renames it using the machine name. No bulk exports.
"""
import argparse
import json
import re
from pathlib import Path
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeout
from test_date_range import desired_range, choose_day, set_clock

ROOT = Path(__file__).resolve().parent
URL = "https://lvmm.justech.com.tw/admin/19/orders"

def filename_safe(s):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", s).strip(" .") or "未命名"

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--month", default="2026-10", help="對帳年月 YYYY-MM")
    args = parser.parse_args()
    names_path = ROOT / "device_names.json"
    if not names_path.exists():
        raise SystemExit("缺少 device_names.json，請先執行 discover_devices.py")
    names = json.loads(names_path.read_text(encoding="utf-8"))
    for idx, name in enumerate(names, 1):
        print(f"{idx:>3}. {name}")
    choice = input("請輸入要測試的設備序號：").strip()
    if not choice.isdigit() or not (1 <= int(choice) <= len(names)):
        raise SystemExit("設備序號不正確")
    name = names[int(choice)-1]
    start, end = desired_range(args.month)
    expected = start.strftime("%d/%m/%Y 12:00:00 AM") + " - " + end.strftime("%d/%m/%Y 12:00:00 AM")

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(ROOT / "browser_profile"), headless=False,
            viewport={"width":1440,"height":900}, accept_downloads=True)
        try:
            page = context.pages[0] if context.pages else context.new_page()
            page.goto(URL, wait_until="domcontentloaded")
            input("請登入 JUSTECH、確認訂單頁已開啟，再按 Enter：")
            if "/login" in page.url:
                page.goto(URL, wait_until="domcontentloaded")
            page.locator(".choices__inner").first.click()
            dropdown = page.locator(".choices__list--dropdown").first
            dropdown.wait_for(state="visible")
            search = dropdown.locator("input.choices__input")
            if search.count():
                search.first.fill(name)
            options = dropdown.locator(".choices__item--choice")
            matches = [i for i in range(options.count()) if " ".join(options.nth(i).inner_text().split()) == name]
            if len(matches) != 1:
                raise RuntimeError(f"設備選項無法唯一匹配 {name}：{len(matches)} 筆")
            options.nth(matches[0]).click()
            field = page.get_by_role("textbox",name="建立時間")
            field.click()
            picker = page.locator(".daterangepicker:visible").first
            choose_day(picker,start)
            choose_day(picker,end)
            set_clock(picker,"left","12","0","0","AM")
            set_clock(picker,"right","12","0","0","AM")
            picker.get_by_role("button",name="套用").click()
            actual = field.input_value().strip()
            if actual != expected:
                raise RuntimeError(f"日期不符：預期 {expected}，實際 {actual}")
            print("已選取設備：", name)
            print("已套用日期：", actual)
            answer = input("請目視核對設備和日期；正確輸入 YES 開始人工匯出／下載，其餘取消：")
            if answer != "YES":
                print("已取消，沒有匯出")
                return
            print("請在瀏覽器手動完成『匯出報表 → 商品月銷售數量統計表 → 確定 → 通知中心 → 下載』。")
            print("最多等待 5 分鐘偵測瀏覽器下載事件。")
            try:
                with page.expect_download(timeout=300000) as download_info:
                    pass  # Downloads happen while event loop waits for the user action.
                download = download_info.value
            except PlaywrightTimeout:
                raise RuntimeError("5 分鐘內未在原訂單分頁偵測到下載，可能需適配通知中心或新分頁")
            # JUSTECH sometimes returns a UUID filename without any extension.
            # Inspect actual bytes before deciding the extension.
            import tempfile
            import zipfile
            import csv
            with tempfile.TemporaryDirectory() as tmpdir:
                temp = Path(tmpdir) / "downloaded_report"
                download.save_as(temp)
                if temp.stat().st_size == 0:
                    raise RuntimeError("下載內容為空")
                with temp.open("rb") as file_handle:
                    header = file_handle.read(8)
                if zipfile.is_zipfile(temp):
                    from openpyxl import load_workbook
                    try:
                        with temp.open("rb") as file_handle:
                            book = load_workbook(file_handle, read_only=True)
                            try:
                                if not book.sheetnames:
                                    raise ValueError("Excel 沒有工作表")
                            finally:
                                book.close()
                    except Exception as exc:
                        raise RuntimeError("下載檔是 ZIP，但無法辨識成 Excel xlsx") from exc
                    extension = ".xlsx"
                elif header.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                    extension = ".xls"
                else:
                    try:
                        sample = temp.read_bytes()[:4096].decode("utf-8-sig")
                        if not sample.strip() or sample.lstrip().lower().startswith(("<html", "<!doctype")):
                            raise ValueError("不是 CSV 內容")
                        rows = list(csv.reader(sample.splitlines()))
                        if len(rows) == 0 or max(map(len, rows)) < 2:
                            raise ValueError("未辨識出 CSV 欄位")
                        extension = ".csv"
                    except Exception as exc:
                        raise RuntimeError(
                            f"檔案格式無法辨識（原始檔名 {download.suggested_filename}），"
                            "先保留人工檢查，不套用 Excel 副檔名"
                        ) from exc
                target_dir = ROOT / "downloads" / args.month
                target_dir.mkdir(parents=True, exist_ok=True)
                target = target_dir / f"{filename_safe(args.month)}_{filename_safe(name)}{extension}"
                if target.exists():
                    raise RuntimeError(f"檔案已存在，為避免覆蓋請先檢查：{target}")
                target.write_bytes(temp.read_bytes())
            print("下載完成：",target)
            print("注意：此為人工確認的單台測試，通知中心任務尚未自動關聯；請核對內容。")
        finally:
            context.close()

if __name__ == "__main__":
    main()
